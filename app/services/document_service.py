from __future__ import annotations

import io
import json
import re
from collections.abc import Mapping
from dataclasses import asdict, dataclass, field
from pathlib import Path
from typing import Any, Iterable

from docx import Document
from docx.enum.text import WD_BREAK
from PIL import Image, ImageOps

from app.image_processing.preprocessing import DocumentNormalizer


IMAGE_SUFFIXES = {".jpg", ".jpeg", ".png", ".tif", ".tiff", ".webp"}


class DocumentImportError(ValueError):
    pass


@dataclass(frozen=True)
class PreparedPage:
    source_name: str
    page_number: int
    filename: str
    content_type: str
    data: bytes


@dataclass(frozen=True)
class ExportPage:
    page_number: int
    filename: str
    text: str
    recognition_id: int | None = None
    model_name: str | None = None
    recognition_mode: str | None = None
    processing_time_ms: int | None = None
    needs_review: bool = False
    line_coordinates: list[dict[str, int]] = field(default_factory=list)


def expand_document(
    filename: str,
    data: bytes,
    content_type: str | None = None,
    *,
    pdf_render_dpi: int = 160,
    max_pages: int = 100,
    max_file_size_bytes: int | None = None,
    max_image_pixels: int = 60_000_000,
) -> list[PreparedPage]:
    if max_file_size_bytes is not None and len(data) > max_file_size_bytes:
        raise DocumentImportError(
            f"Файл {filename} превышает допустимый размер {max_file_size_bytes // (1024 * 1024)} МБ."
        )
    suffix = Path(filename).suffix.lower()
    if suffix == ".pdf" or content_type == "application/pdf":
        return _render_pdf_pages(
            filename,
            data,
            pdf_render_dpi=pdf_render_dpi,
            max_pages=max_pages,
        )
    if suffix not in IMAGE_SUFFIXES:
        raise DocumentImportError(f"Неподдерживаемый формат файла: {suffix or 'без расширения'}.")
    if not data:
        raise DocumentImportError(f"Файл {filename} пуст.")
    data, normalized_suffix = _normalize_exif_orientation(
        data,
        suffix,
        filename=filename,
        max_image_pixels=max_image_pixels,
    )
    if normalized_suffix != suffix:
        filename = str(Path(filename).with_suffix(normalized_suffix))
        suffix = normalized_suffix
    return [
        PreparedPage(
            source_name=filename,
            page_number=1,
            filename=filename,
            content_type=content_type or _image_content_type(suffix),
            data=data,
        )
    ]


def build_export_pages(results: Iterable[dict[str, Any]]) -> list[ExportPage]:
    pages = []
    for index, result in enumerate(results, start=1):
        pages.append(
            ExportPage(
                page_number=int(result.get("document_page_number") or result.get("page_number") or index),
                filename=str(result.get("filename") or result.get("original_filename") or f"page_{index}.png"),
                text=str(
                    result.get("edited_text")
                    or result.get("text")
                    or result.get("corrected_text")
                    or result.get("recognized_text")
                    or ""
                ),
                recognition_id=result.get("id"),
                model_name=result.get("model_name"),
                recognition_mode=result.get("recognition_mode"),
                processing_time_ms=result.get("processing_time_ms"),
                needs_review=bool(result.get("needs_review", False)),
                line_coordinates=list(result.get("line_coordinates") or []),
            )
        )
    return pages


def load_image_folder(
    folder: str | Path,
    *,
    recursive: bool = False,
    max_pages: int = 500,
    max_file_size_bytes: int | None = None,
    max_image_pixels: int = 60_000_000,
    sort_by: str = "name",
) -> list[PreparedPage]:
    source = Path(folder).expanduser().resolve()
    if not source.is_dir():
        raise DocumentImportError(f"Папка не найдена: {source}")

    iterator = source.rglob("*") if recursive else source.iterdir()
    if sort_by not in {"name", "modified", "created"}:
        raise DocumentImportError(f"Неизвестный порядок сортировки: {sort_by}")
    candidates = [
        path for path in iterator if path.is_file() and path.suffix.lower() in IMAGE_SUFFIXES
    ]
    if sort_by == "name":
        image_paths = sorted(
            candidates,
            key=lambda path: _natural_key(str(path.relative_to(source))),
        )
    else:
        timestamp_name = "st_mtime" if sort_by == "modified" else "st_ctime"
        image_paths = sorted(
            candidates,
            key=lambda path: (
                getattr(path.stat(), timestamp_name),
                _natural_key(str(path.relative_to(source))),
            ),
        )
    if not image_paths:
        raise DocumentImportError("В выбранной папке нет поддерживаемых изображений.")
    if len(image_paths) > max_pages:
        raise DocumentImportError(
            f"В папке найдено {len(image_paths)} изображений; разрешено не более {max_pages}."
        )

    pages: list[PreparedPage] = []
    for index, path in enumerate(image_paths, start=1):
        prepared = expand_document(
            path.name,
            path.read_bytes(),
            _image_content_type(path.suffix.lower()),
            max_pages=1,
            max_file_size_bytes=max_file_size_bytes,
            max_image_pixels=max_image_pixels,
        )[0]
        pages.append(
            PreparedPage(
                source_name=str(path.relative_to(source)),
                page_number=index,
                filename=prepared.filename,
                content_type=prepared.content_type,
                data=prepared.data,
            )
        )
    return pages


def transform_image(
    data: bytes,
    *,
    rotation_degrees: int = 0,
    crop_box: tuple[int, int, int, int] | None = None,
    correct_perspective: bool = False,
    deskew: bool = False,
    enhance_contrast: bool = False,
    denoise: bool = False,
    remove_grid: bool = False,
) -> bytes:
    if not data:
        raise DocumentImportError("Нельзя преобразовать пустое изображение.")

    try:
        image = ImageOps.exif_transpose(Image.open(io.BytesIO(data))).convert("RGB")
    except Exception as exc:
        raise DocumentImportError("Не удалось открыть изображение для преобразования.") from exc

    normalized_rotation = rotation_degrees % 360
    if normalized_rotation not in {0, 90, 180, 270}:
        raise DocumentImportError("Поворот должен быть кратен 90 градусам.")
    if normalized_rotation:
        image = image.rotate(-normalized_rotation, expand=True, fillcolor="white")

    if crop_box is not None:
        left, top, right, bottom = crop_box
        if left < 0 or top < 0 or right > image.width or bottom > image.height:
            raise DocumentImportError("Область обрезки выходит за границы изображения.")
        if right <= left or bottom <= top:
            raise DocumentImportError("Область обрезки должна иметь положительный размер.")
        image = image.crop(crop_box)

    if correct_perspective or deskew or enhance_contrast or denoise or remove_grid:
        import cv2
        import numpy as np

        array = cv2.cvtColor(np.asarray(image), cv2.COLOR_RGB2BGR)
        if correct_perspective:
            array = DocumentNormalizer().normalize(array)
        gray = cv2.cvtColor(array, cv2.COLOR_BGR2GRAY)
        if deskew:
            gray = _deskew_image(gray)
        if denoise:
            gray = cv2.bilateralFilter(gray, d=7, sigmaColor=50, sigmaSpace=50)
        if enhance_contrast:
            gray = cv2.createCLAHE(clipLimit=2.0, tileGridSize=(8, 8)).apply(gray)
        if remove_grid:
            gray = _remove_notebook_grid(gray)
        image = Image.fromarray(gray)

    output = io.BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def export_text(pages: Iterable[ExportPage], *, page_headers: bool = True) -> str:
    blocks = []
    for page in pages:
        text = page.text.strip()
        if page_headers:
            blocks.append(f"Страница {page.page_number}\n\n{text}".rstrip())
        else:
            blocks.append(text)
    return "\n\n\f\n\n".join(blocks).strip()


def export_json(pages: Iterable[ExportPage], *, title: str = "Распознанный документ") -> bytes:
    payload = {
        "title": title,
        "pages": [asdict(page) for page in pages],
    }
    return json.dumps(payload, ensure_ascii=False, indent=2).encode("utf-8")


def export_docx(pages: Iterable[ExportPage], *, title: str = "Распознанный документ") -> bytes:
    page_list = list(pages)
    document = Document()
    document.core_properties.title = title
    document.add_heading(title, level=0)

    for index, page in enumerate(page_list):
        document.add_heading(f"Страница {page.page_number}", level=1)
        for line in page.text.splitlines() or [""]:
            document.add_paragraph(line)
        if index + 1 < len(page_list):
            document.paragraphs[-1].add_run().add_break(WD_BREAK.PAGE)

    output = io.BytesIO()
    document.save(output)
    return output.getvalue()


def export_searchable_pdf(
    pages: Iterable[ExportPage],
    page_images: Mapping[int, bytes],
) -> bytes:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentImportError(
            "Для создания PDF установите зависимость pymupdf."
        ) from exc

    document = fitz.open()
    font_path = _searchable_pdf_font()
    try:
        for export_page in pages:
            image_data = page_images.get(export_page.page_number)
            if not image_data:
                raise DocumentImportError(
                    f"Не найдено изображение для страницы {export_page.page_number}."
                )
            try:
                with Image.open(io.BytesIO(image_data)) as image:
                    width, height = image.size
            except Exception as exc:
                raise DocumentImportError(
                    f"Не удалось прочитать изображение страницы {export_page.page_number}."
                ) from exc

            pdf_page = document.new_page(width=width, height=height)
            pdf_page.insert_image(pdf_page.rect, stream=image_data)
            font_name = "ocr-font"
            if font_path is not None:
                pdf_page.insert_font(fontname=font_name, fontfile=str(font_path))
            else:
                font_name = "helv"
            _insert_searchable_text(pdf_page, export_page, font_name=font_name)

        return document.tobytes(garbage=4, deflate=True)
    finally:
        document.close()


def _render_pdf_pages(
    filename: str,
    data: bytes,
    *,
    pdf_render_dpi: int,
    max_pages: int,
) -> list[PreparedPage]:
    try:
        import fitz
    except ImportError as exc:
        raise DocumentImportError(
            "Для открытия PDF установите зависимость pymupdf."
        ) from exc

    try:
        document = fitz.open(stream=data, filetype="pdf")
    except Exception as exc:
        raise DocumentImportError(f"Не удалось открыть PDF {filename}.") from exc

    try:
        if document.page_count == 0:
            raise DocumentImportError(f"PDF {filename} не содержит страниц.")
        if document.page_count > max_pages:
            raise DocumentImportError(
                f"PDF содержит {document.page_count} страниц; разрешено не более {max_pages}."
            )
        scale = max(72, pdf_render_dpi) / 72
        matrix = fitz.Matrix(scale, scale)
        stem = Path(filename).stem
        return [
            PreparedPage(
                source_name=filename,
                page_number=index + 1,
                filename=f"{stem}_page_{index + 1:03d}.png",
                content_type="image/png",
                data=document.load_page(index).get_pixmap(matrix=matrix, alpha=False).tobytes("png"),
            )
            for index in range(document.page_count)
        ]
    finally:
        document.close()


def _image_content_type(suffix: str) -> str:
    return {
        ".jpg": "image/jpeg",
        ".jpeg": "image/jpeg",
        ".png": "image/png",
        ".tif": "image/tiff",
        ".tiff": "image/tiff",
        ".webp": "image/webp",
    }.get(suffix, "application/octet-stream")


def _normalize_exif_orientation(
    data: bytes,
    suffix: str,
    *,
    filename: str,
    max_image_pixels: int,
) -> tuple[bytes, str]:
    try:
        with Image.open(io.BytesIO(data)) as source:
            if source.width * source.height > max_image_pixels:
                raise DocumentImportError(
                    f"Изображение {filename} содержит слишком много пикселей: "
                    f"{source.width}x{source.height}."
                )
            orientation = source.getexif().get(274, 1)
            if orientation == 1:
                return data, suffix
            normalized = ImageOps.exif_transpose(source).convert("RGB")
            output = io.BytesIO()
            normalized.save(output, format="PNG", optimize=True)
            return output.getvalue(), ".png"
    except DocumentImportError:
        raise
    except Exception as exc:
        raise DocumentImportError(f"Не удалось открыть изображение {filename}.") from exc


def _deskew_image(gray: Any) -> Any:
    import cv2
    import numpy as np

    _, foreground = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)
    coordinates = np.column_stack(np.where(foreground > 0))
    if len(coordinates) < 20:
        return gray
    angle = cv2.minAreaRect(coordinates[:, ::-1].astype(np.float32))[-1]
    if angle > 45:
        angle -= 90
    if abs(angle) < 0.15 or abs(angle) > 15:
        return gray
    height, width = gray.shape
    matrix = cv2.getRotationMatrix2D((width / 2, height / 2), angle, 1.0)
    return cv2.warpAffine(
        gray,
        matrix,
        (width, height),
        flags=cv2.INTER_CUBIC,
        borderMode=cv2.BORDER_CONSTANT,
        borderValue=255,
    )


def _remove_notebook_grid(gray: Any) -> Any:
    import cv2
    import numpy as np

    _, ink = cv2.threshold(gray, 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)
    horizontal_size = max(12, gray.shape[1] // 24)
    vertical_size = max(12, gray.shape[0] // 24)
    horizontal = cv2.morphologyEx(
        ink,
        cv2.MORPH_OPEN,
        np.ones((1, horizontal_size), dtype=np.uint8),
    )
    vertical = cv2.morphologyEx(
        ink,
        cv2.MORPH_OPEN,
        np.ones((vertical_size, 1), dtype=np.uint8),
    )
    grid = cv2.bitwise_or(horizontal, vertical)
    cleaned = gray.copy()
    cleaned[grid > 0] = 255
    return cleaned


def _insert_searchable_text(page: Any, export_page: ExportPage, *, font_name: str) -> None:
    lines = export_page.text.splitlines()
    coordinates = export_page.line_coordinates
    if coordinates and lines:
        for text, box in zip(lines, coordinates, strict=False):
            if not text.strip():
                continue
            box_width = max(1, box["x2"] - box["x1"])
            box_height = max(1, box["y2"] - box["y1"])
            width_limited_size = box_width / max(1, len(text)) * 1.7
            font_size = max(3.0, min(24.0, box_height * 0.72, width_limited_size))
            page.insert_text(
                (box["x1"], box["y2"] - max(1.0, box_height * 0.08)),
                text,
                fontname=font_name,
                fontsize=font_size,
                render_mode=3,
                overlay=True,
            )
        return

    if export_page.text.strip():
        margin = max(8.0, min(page.rect.width, page.rect.height) * 0.02)
        font_size = max(4.0, min(12.0, page.rect.height / 60))
        y = margin + font_size
        for text in export_page.text.splitlines():
            if text.strip():
                page.insert_text(
                    (margin, y),
                    text,
                    fontname=font_name,
                    fontsize=font_size,
                    render_mode=3,
                    overlay=True,
                )
            y += font_size * 1.35
            if y >= page.rect.height - margin:
                break


def _natural_key(value: str) -> list[str | int]:
    return [int(part) if part.isdigit() else part.casefold() for part in re.split(r"(\d+)", value)]


def _searchable_pdf_font() -> Path | None:
    candidates = (
        Path("C:/Windows/Fonts/arial.ttf"),
        Path("C:/Windows/Fonts/calibri.ttf"),
        Path("/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf"),
        Path("/usr/share/fonts/truetype/liberation2/LiberationSans-Regular.ttf"),
    )
    return next((path for path in candidates if path.is_file()), None)
